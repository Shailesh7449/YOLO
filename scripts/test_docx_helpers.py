import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.0)

# Test styled table
tbl = doc.add_table(rows=2, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = tbl.rows[0].cells
hdr_cells[0].text = "Header 1"
hdr_cells[1].text = "Header 2"

for c in hdr_cells:
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
    c._tc.get_or_add_tcPr().append(shd)
    for p in c.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.bold = True

doc.save("test_styled.docx")
print("Saved test_styled.docx successfully!")
