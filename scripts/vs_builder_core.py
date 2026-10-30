import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT_PATH = r"d:\spidy\YOLO\VisionSense_BCSE497J_Project_I_Report.docx"
FIGURES_DIR = r"d:\spidy\YOLO\report_figures"

def sanitize_text(text: str) -> str:
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="1E3A8A"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

class VisionSenseReportBuilder:
    def __init__(self):
        self.doc = Document()
        self._configure_styles()

    def _configure_styles(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)

    def add_page_break(self):
        self.doc.add_page_break()

    def add_chapter_title(self, number_str: str, title_str: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"CHAPTER {number_str}: {title_str.upper()}")
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_frontmatter_title(self, title_str: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title_str.upper())
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def add_section_title(self, number_str: str, title_str: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{number_str} {title_str}")
        run.bold = True
        run.font.size = Pt(13.5)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def add_subsection_title(self, number_str: str, title_str: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{number_str} {title_str}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    def add_body_p(self, text: str, italic=False, bold=False, indent=False, align_justify=True):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        if align_justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.3)
        run = p.add_run(sanitize_text(text))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.italic = italic
        run.bold = bold
        return p

    def add_bullet_point(self, prefix: str, text: str):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_p = p.add_run(sanitize_text(prefix))
        run_p.bold = True
        run_p.font.name = 'Times New Roman'
        run_p.font.size = Pt(11.5)
        
        run_t = p.add_run(" " + sanitize_text(text))
        run_t.font.name = 'Times New Roman'
        run_t.font.size = Pt(11.5)
        return p

    def add_equation_box(self, eq_text: str, eq_num: str = ""):
        tbl = self.doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="E5E7EB", sz="2")
        
        row = tbl.rows[0]
        prevent_row_split(row)
        c_left = row.cells[0]
        c_right = row.cells[1]
        
        c_left.width = Inches(5.3)
        c_right.width = Inches(0.9)
        
        set_cell_margins(c_left, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c_right, top=100, bottom=100, left=140, right=140)
        set_cell_shading(c_left, "F8FAFC")
        set_cell_shading(c_right, "F8FAFC")
        
        p_eq = c_left.paragraphs[0]
        p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_eq = p_eq.add_run(sanitize_text(eq_text))
        r_eq.font.name = 'Times New Roman'
        r_eq.font.size = Pt(11.5)
        r_eq.italic = True
        
        p_num = c_right.paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_num = p_num.add_run(f"({eq_num})" if eq_num else "")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(11)
        r_num.bold = True

    def add_table_custom(self, headers: list, data: list, col_widths: list, caption_title: str, table_num: str):
        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(12)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_num = p_cap.add_run(f"Table {table_num}: ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_title = p_cap.add_run(sanitize_text(caption_title))
        r_title.bold = True
        r_title.italic = True
        r_title.font.size = Pt(10)
        
        num_cols = len(headers)
        table = self.doc.add_table(rows=len(data) + 1, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="D1D5DB", sz="4")
        
        hdr_row = table.rows[0]
        make_row_header(hdr_row)
        for i, h_text in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = Inches(col_widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
            set_cell_shading(cell, "1E3A8A")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(sanitize_text(h_text))
            r.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        for r_idx, row_data in enumerate(data):
            row = table.rows[r_idx + 1]
            prevent_row_split(row)
            bg_color = "F9FAFB" if (r_idx % 2 == 1) else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = Inches(col_widths[c_idx])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                set_cell_shading(cell, bg_color)
                
                p = cell.paragraphs[0]
                if len(str(val)) <= 10 or "%" in str(val) or str(val).replace('.', '', 1).isdigit():
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(sanitize_text(str(val)))
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                if "**" in str(val):
                    r.bold = True
                    r.text = r.text.replace("**", "")

        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(6)

    def add_figure_image(self, img_name: str, fig_num: str, caption_title: str, width_in=5.8):
        img_path = os.path.join(FIGURES_DIR, img_name)
        if not os.path.exists(img_path):
            print(f"Warning: image {img_path} not found!")
            return
            
        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))
        
        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_num = p_cap.add_run(f"Figure {fig_num}: ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_title = p_cap.add_run(sanitize_text(caption_title))
        r_title.bold = True
        r_title.italic = True
        r_title.font.size = Pt(10)

    def add_code_block(self, code_text: str, filename: str = ""):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="CBD5E1", sz="4")
        
        row = tbl.rows[0]
        prevent_row_split(row)
        cell = row.cells[0]
        cell.width = Inches(6.2)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_shading(cell, "F1F5F9")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        if filename:
            r_fn = p.add_run(f"// {filename}\n")
            r_fn.bold = True
            r_fn.font.name = 'Consolas'
            r_fn.font.size = Pt(8.5)
            r_fn.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
            
        r_code = p.add_run(sanitize_text(code_text))
        r_code.font.name = 'Consolas'
        r_code.font.size = Pt(8.5)
        r_code.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        p_sp = self.doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(6)

print("VisionSense builder core initialized.")
