import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()
sec0 = doc.sections[0]
sec0.top_margin = Inches(1.0)
sec0.bottom_margin = Inches(1.0)
sec0.left_margin = Inches(1.25)
sec0.right_margin = Inches(1.0)
doc.add_paragraph("Cover Page - No Page Number")

# Section 1: Front Matter (Roman)
sec1 = doc.add_section()
sec1.header.is_linked_to_previous = False
sec1.footer.is_linked_to_previous = False

sectPr1 = sec1._sectPr
pgNumType1 = sectPr1.find(qn('w:pgNumType'))
if pgNumType1 is None:
    pgNumType1 = OxmlElement('w:pgNumType')
    sectPr1.append(pgNumType1)
pgNumType1.set(qn('w:fmt'), 'lowerRoman')
pgNumType1.set(qn('w:start'), '1')

p1 = sec1.footer.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p1.add_run()
r1._r.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))
doc.add_paragraph("Front Matter Page 1 (Declaration)")

# Section 2: Main Body (Decimal)
sec2 = doc.add_section()
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False

sectPr2 = sec2._sectPr
pgNumType2 = sectPr2.find(qn('w:pgNumType'))
if pgNumType2 is None:
    pgNumType2 = OxmlElement('w:pgNumType')
    sectPr2.append(pgNumType2)
pgNumType2.set(qn('w:fmt'), 'decimal')
pgNumType2.set(qn('w:start'), '1')

p_hdr = sec2.header.paragraphs[0]
p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_hdr = p_hdr.add_run("VIT BCSE497J Project-I | Multimodal Meeting Auditor")
r_hdr.font.size = Pt(8.5)
r_hdr.font.color.rgb = RGBColor(100, 116, 139)

p2 = sec2.footer.paragraphs[0]
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run()
r2._r.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))
doc.add_paragraph("Chapter 1 Body Page 1")

doc.save("test_sections.docx")
print("Saved test_sections.docx successfully!")
